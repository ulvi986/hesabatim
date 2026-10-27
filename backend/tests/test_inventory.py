from decimal import Decimal

from tests.conftest import create_product


def test_create_product_records_initial_event(make_user):
    s = make_user()
    p = create_product(s, quantity="500", purchase_price="60", selling_price="85")
    assert p["quantity"] == "500"
    assert p["inventory_value"] == "30000.00"  # 500 × 60 at cost
    assert p["retail_value"] == "42500.00"
    assert p["stock_status"] == "in_stock"
    assert p["sku"]  # generated
    events = s.get(s.biz(f"/inventory/events?product_id={p['id']}")).json()["items"]
    assert [e["event_type"] for e in events] == ["initial"]


def test_inventory_value_is_never_guessed_from_selling_price(make_user):
    s = make_user()
    p = create_product(s, purchase_price=None, selling_price="85", quantity="10")
    assert p["inventory_value"] is None
    summ = s.get(s.biz("/inventory/summary")).json()
    assert summ["inventory_value"] == "0.00" and summ["products_without_cost"] == 1


def test_product_fields_validation(make_user):
    s = make_user()
    assert s.post(s.biz("/products"), json={"name": "Q"}).status_code == 422
    assert s.post(s.biz("/products"), json={"name": "Qapı", "quantity": "-1"}).status_code == 422
    assert s.post(s.biz("/products"), json={"name": "Qapı", "selling_price": "-5"}).status_code == 422
    assert s.post(s.biz("/products"), json={"name": "Qapı", "unit": "bucket"}).status_code == 422
    create_product(s, sku="DR-1")
    r = s.post(s.biz("/products"), json={"name": "Başqa", "sku": "DR-1"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "sku_taken"


def test_edit_product_and_audit(make_user):
    s = make_user()
    p = create_product(s)
    r = s.patch(s.biz(f"/products/{p['id']}"), json={"selling_price": "95", "description": "Otaqarası"})
    assert r.status_code == 200 and r.json()["selling_price"] == "95.00"
    logs = s.get(s.biz("/audit-logs")).json()["items"]
    assert logs[0]["action"] == "product.update" and logs[0]["before"]["selling_price"] == "85.00"


def test_categories_crud_and_filter(make_user):
    s = make_user()
    cat = s.post(s.biz("/categories"), json={"name": "Qapılar"}).json()
    assert s.post(s.biz("/categories"), json={"name": "qapılar"}).status_code == 409  # case-insensitive unique
    p = create_product(s, category_id=cat["id"])
    create_product(s, name="Masa")
    res = s.get(s.biz(f"/products?category_id={cat['id']}")).json()
    assert [i["id"] for i in res["items"]] == [p["id"]]
    assert s.get(s.biz("/categories")).json()["items"][0]["product_count"] == 1
    assert s.delete(s.biz(f"/categories/{cat['id']}")).status_code == 204
    assert s.get(s.biz(f"/products/{p['id']}")).json()["category"] is None  # product kept, uncategorised


def test_adjust_stock_updates_quantity_and_audits(make_user):
    s = make_user()
    p = create_product(s, quantity="100")
    r = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-40", "reason": "Xarab oldu"})
    assert r.status_code == 200, r.text
    assert r.json()["product"]["quantity"] == "60"
    ev = r.json()["event"]
    assert (ev["quantity_before"], ev["quantity_after"], ev["delta"]) == ("100", "60", "-40")
    logs = s.get(s.biz("/audit-logs")).json()["items"]
    adj = [l for l in logs if l["action"] == "inventory.adjust"]
    assert adj and adj[0]["after"]["quantity"] == "60.000"


def test_stock_cannot_go_negative(make_user):
    s = make_user()
    p = create_product(s, quantity="10")
    r = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-11"})
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "insufficient_stock"
    assert s.get(s.biz(f"/products/{p['id']}")).json()["quantity"] == "10"


def test_set_quantity_and_fractional_units(make_user):
    s = make_user()
    p = create_product(s, name="Laminat 8mm", unit="m2", quantity="120.5")
    r = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"set_quantity": "99.25"})
    assert r.json()["product"]["quantity"] == "99.25"


def test_idempotent_adjust(make_user):
    s = make_user()
    p = create_product(s, quantity="100")
    hdr = {"Idempotency-Key": "adj-123"}
    r1 = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-5"}, headers=hdr)
    r2 = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-5"}, headers=hdr)
    assert r1.json() == r2.json()
    assert s.get(s.biz(f"/products/{p['id']}")).json()["quantity"] == "95"


def test_optimistic_version_conflict(make_user):
    s = make_user()
    p = create_product(s, quantity="100")
    v = p["inventory_version"]
    assert s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "1", "expected_version": v}).status_code == 200
    r = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "1", "expected_version": v})
    assert r.status_code == 409 and r.json()["error"]["code"] == "stale_version"


def test_undo_latest_event_only(make_user):
    s = make_user()
    p = create_product(s, quantity="100")
    e1 = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-10"}).json()["event"]
    e2 = s.post(s.biz(f"/inventory/{p['id']}/adjust"), json={"delta": "-20"}).json()["event"]
    listed = {e["id"]: e for e in s.get(s.biz("/inventory/events")).json()["items"]}
    assert listed[e2["id"]]["undoable"] and not listed[e1["id"]]["undoable"]
    assert s.post(s.biz(f"/inventory/events/{e1['id']}/undo")).status_code == 409
    r = s.post(s.biz(f"/inventory/events/{e2['id']}/undo"))
    assert r.status_code == 200
    assert r.json()["product"]["quantity"] == "90"
    assert s.post(s.biz(f"/inventory/events/{e2['id']}/undo")).status_code == 409  # not twice


def test_archive_and_restore(make_user):
    s = make_user()
    p = create_product(s)
    assert s.post(s.biz("/products/archive"), json={"ids": [p["id"]]}).json()["updated"] == 1
    assert s.get(s.biz("/products")).json()["total"] == 0
    assert s.get(s.biz("/products?archived=true")).json()["total"] == 1
    s.post(s.biz("/products/archive"), json={"ids": [p["id"]], "archived": False})
    assert s.get(s.biz("/products")).json()["total"] == 1


def test_low_stock_filter_and_summary(make_user):
    s = make_user()
    create_product(s, name="Ağ MDF qapı", quantity="5", low_stock_threshold=10, purchase_price="50")
    create_product(s, name="Taxta lövhə", quantity="300", purchase_price="10")
    create_product(s, name="Kafel 30x30", quantity="0", unit="m2", purchase_price="20")
    low = s.get(s.biz("/products?stock=attention")).json()
    assert {i["name"] for i in low["items"]} == {"Ağ MDF qapı", "Kafel 30x30"}
    summ = s.get(s.biz("/inventory/summary")).json()
    assert summ["total_products"] == 3
    assert Decimal(summ["inventory_value"]) == Decimal("3250.00")  # 5×50 + 300×10
    assert summ["out_of_stock_count"] == 1 and summ["low_stock_count"] == 2


def test_search_handles_azerbaijani_forms(make_user):
    s = make_user()
    create_product(s, name="Ağ MDF qapı")
    create_product(s, name="Taxta lövhə")
    create_product(s, name="Masa")
    res = s.get(s.biz("/products"), params={"q": "qapılar"}).json()
    assert [i["name"] for i in res["items"]] == ["Ağ MDF qapı"]
    res = s.get(s.biz("/products"), params={"q": "white door"}).json()
    assert [i["name"] for i in res["items"]] == ["Ağ MDF qapı"]
    res = s.get(s.biz("/products"), params={"q": "masanı"}).json()
    assert [i["name"] for i in res["items"]] == ["Masa"]
    p = s.get(s.biz("/products"), params={"q": "Masa"}).json()["items"][0]
    assert [i["name"] for i in s.get(s.biz("/products"), params={"q": p["sku"]}).json()["items"]] == ["Masa"]


def test_csv_import_and_export(make_user):
    s = make_user()
    csv_content = ("sku,name,quantity,unit,purchase_price,selling_price,category\n"
                   "DR-1,Palıd qapı,40,ədəd,120,180,Qapılar\n,Gips 25kg,abc,kq,5,7,\n")
    files = {"file": ("inv.csv", csv_content.encode("utf-8"), "text/csv")}
    dry = s.post(s.biz("/inventory/import?dry_run=true"), files=files).json()
    assert dry["created"] == 1 and dry["errors"] == 1
    assert s.get(s.biz("/products")).json()["total"] == 0  # dry run wrote nothing
    assert s.get(s.biz("/categories")).json()["items"] == []
    real = s.post(s.biz("/inventory/import?dry_run=false"), files=files).json()
    assert real["created"] == 1
    item = s.get(s.biz("/products")).json()["items"][0]
    assert item["category"]["name"] == "Qapılar"
    # re-import updates quantity via a ledger event instead of overwriting silently
    csv2 = "sku,name,quantity\nDR-1,Palıd qapı,35\n"
    s.post(s.biz("/inventory/import?dry_run=false"), files={"file": ("i.csv", csv2.encode(), "text/csv")})
    items = s.get(s.biz("/products")).json()["items"]
    assert items[0]["quantity"] == "35"
    events = s.get(s.biz(f"/inventory/events?product_id={items[0]['id']}")).json()["items"]
    assert events[0]["event_type"] == "import" and events[0]["delta"] == "-5"
    exp = s.get(s.biz("/inventory/export?format=csv"))
    assert exp.status_code == 200 and "Palıd qapı" in exp.content.decode("utf-8-sig")
    assert s.get(s.biz("/inventory/export?format=xlsx")).status_code == 200


def test_xlsx_import(make_user):
    import io

    from openpyxl import Workbook

    s = make_user()
    wb = Workbook()
    wb.active.append(["Ad", "Miqdar", "Satış qiyməti"])
    wb.active.append(["Stul", 12, 45])
    buf = io.BytesIO()
    wb.save(buf)
    r = s.post(s.biz("/inventory/import?dry_run=false"),
               files={"file": ("inv.xlsx", buf.getvalue(), "application/octet-stream")}).json()
    assert r["created"] == 1
    p = s.get(s.biz("/products")).json()["items"][0]
    assert (p["name"], p["quantity"], p["selling_price"]) == ("Stul", "12", "45.00")


def test_import_rejects_unknown_format(make_user):
    s = make_user()
    r = s.post(s.biz("/inventory/import"), files={"file": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 422


def _xlsx(*rows, title=None):
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if title:
        ws.title = title
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_import_finds_header_under_title_and_cleans_numbers(make_user):
    s = make_user()
    data = _xlsx(["Anbar qalığı — oktyabr"], [], ["№", "Məhsulun adı", "Qalıq", "Ölçü vahidi", "Satış qiyməti (AZN)", "Rəng"],
                 [1, "Palıd qapı", 40, "ədəd", "1 250,50 AZN", "qəhvəyi"], [2, "Laminat", "12,5", "m2", "18", ""],
                 [None, "Cəmi", 52.5, None, None, None])
    dry = s.post(s.biz("/inventory/import?dry_run=true"), files={"file": ("anbar.xlsx", data, "application/octet-stream")}).json()
    assert dry["created"] == 2 and dry["errors"] == 0 and dry["guessed"] is False
    assert [c["field"] for c in dry["columns"]] == ["name", "quantity", "unit", "selling_price"]
    assert dry["ignored"] == ["№", "Rəng"]
    first = dry["rows"][0]
    assert (first["row"], first["name"], first["quantity"], first["selling_price"]) == (4, "Palıd qapı", "40", "1250.50")
    s.post(s.biz("/inventory/import?dry_run=false"), files={"file": ("anbar.xlsx", data, "application/octet-stream")})
    items = {p["name"]: p for p in s.get(s.biz("/products")).json()["items"]}
    assert items["Laminat"]["quantity"] == "12.5" and items["Laminat"]["unit"] == "m2"
    assert items["Palıd qapı"]["selling_price"] == "1250.50"

    again = s.post(s.biz("/inventory/import?dry_run=true"),
                   files={"file": ("a.xlsx", _xlsx(["Наименование", "Количество"], ["Laminat", 20]), "x")}).json()
    assert again["updated"] == 1 and again["rows"][0]["old_quantity"] == "12.5" and again["rows"][0]["quantity"] == "20"


def test_xlsx_import_without_header_and_on_second_sheet(make_user):
    s = make_user()
    r = s.post(s.biz("/inventory/import?dry_run=true"),
               files={"file": ("x.xlsx", _xlsx(["Stul", 12], ["Masa", 3]), "x")}).json()
    assert r["guessed"] is True and r["created"] == 2 and r["rows"][0]["name"] == "Stul"

    import io

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Qeydlər"])
    ws = wb.create_sheet("Mallar")
    ws.append(["Mal", "Say", "Qiymət"])
    ws.append(["Dolab", 4, 300])
    buf = io.BytesIO()
    wb.save(buf)
    r = s.post(s.biz("/inventory/import?dry_run=true"), files={"file": ("x.xlsx", buf.getvalue(), "x")}).json()
    assert r["created"] == 1 and r["rows"][0]["name"] == "Dolab" and r["rows"][0]["selling_price"] == "300"


def test_import_rejects_old_xls_and_broken_files(make_user):
    s = make_user()
    r = s.post(s.biz("/inventory/import"), files={"file": ("x.xls", b"\xd0\xcf\x11\xe0", "x")})
    assert r.status_code == 422 and ".xlsx" in r.json()["error"]["message"]
    assert s.post(s.biz("/inventory/import"), files={"file": ("x.xlsx", b"not a zip", "x")}).status_code == 422


def test_csv_with_title_line_semicolons_and_decimal_commas(make_user):
    s = make_user()
    text = ("Anbar qalığı — oktyabr\n\nMəhsulun adı;Qalıq;Ölçü vahidi;Alış qiyməti;Satış qiyməti (AZN);Kateqoriya\n"
            "Ağ laminat qapı;25;ədəd;95,50;149,90 AZN;Qapılar\nPalıd parket;120,5;m2;18;27,5;Döşəmə\n"
            "Şkaf dəstəyi;abc;ədəd;2;4;Aksesuar\nCəmi;145;;;;\n")
    r = s.post(s.biz("/inventory/import?dry_run=true"), files={"file": ("a.csv", text.encode(), "text/csv")}).json()
    assert (r["created"], r["errors"], r["guessed"]) == (2, 1, False)
    assert [x["name"] for x in r["rows"]] == ["Ağ laminat qapı", "Palıd parket", "Şkaf dəstəyi"]
    assert r["rows"][0]["purchase_price"] == "95.50" and r["rows"][1]["quantity"] == "120.5"
