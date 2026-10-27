"""Stock ledger. Every quantity change is a row in inventory_events, executed under a row lock.

Mutating functions flush but never commit; callers own the transaction."""

from __future__ import annotations

import csv
import io
import re
import uuid
from decimal import Decimal
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.core.errors import Conflict, NotFound, ValidationFailed
from app.models import Inventory, InventoryEvent, Product
from app.services.common import audit, iso, money_str, qty, qty_str

NOT_UNDOABLE = ("undo", "sale", "return")


def _lock_inventory(db: Session, business_id: uuid.UUID, product_id: uuid.UUID) -> Inventory:
    inv = db.scalar(
        select(Inventory)
        .where(Inventory.business_id == business_id, Inventory.product_id == product_id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    if inv is None:
        raise NotFound("Məhsul tapılmadı")
    return inv


def adjust_stock(
    db: Session,
    *,
    actor_id: uuid.UUID | None,
    business_id: uuid.UUID,
    product_id: uuid.UUID,
    delta: Any = None,
    set_to: Any = None,
    reason: str | None = None,
    source: str = "manual",
    event_type: str | None = None,
    ai_action_id: uuid.UUID | None = None,
    sale_id: uuid.UUID | None = None,
    idempotency_key: str | None = None,
    expected_version: int | None = None,
) -> InventoryEvent:
    if (delta is None) == (set_to is None):
        raise ValidationFailed("Ya dəyişiklik (delta), ya da yeni miqdar göstərilməlidir")
    if idempotency_key:
        prior = db.scalar(
            select(InventoryEvent).where(InventoryEvent.business_id == business_id,
                                         InventoryEvent.idempotency_key == idempotency_key)
        )
        if prior:
            return prior
    inv = _lock_inventory(db, business_id, product_id)
    if expected_version is not None and inv.version != expected_version:
        raise Conflict("Bu məhsulun stoku başqa istifadəçi tərəfindən dəyişdirilib. Səhifəni yeniləyin.",
                       code="stale_version", details={"current_version": inv.version})
    before = Decimal(inv.quantity)
    if set_to is not None:
        after = qty(set_to)
        change = after - before
        etype = event_type or "set"
    else:
        change = qty(delta)
        after = before + change
        etype = event_type or ("increase" if change >= 0 else "decrease")
    if after < 0:
        raise ValidationFailed(
            f"Stok kifayət etmir: mövcud {qty_str(before)}, tələb olunan azalma {qty_str(-change)}",
            code="insufficient_stock",
            details={"available": qty_str(before), "requested": qty_str(-change)},
        )
    if change == 0:
        raise ValidationFailed("Miqdar dəyişmir", code="no_change")
    inv.quantity = after
    inv.version += 1
    event = InventoryEvent(
        business_id=business_id, product_id=product_id, inventory_id=inv.id, event_type=etype, source=source,
        delta=change, quantity_before=before, quantity_after=after, reason=(reason or None) and reason[:500],
        actor_id=actor_id, ai_action_id=ai_action_id, sale_id=sale_id, idempotency_key=idempotency_key,
    )
    db.add(event)
    audit(db, business_id=business_id, actor_id=actor_id, action="inventory.adjust", entity_type="product",
          entity_id=product_id, before={"quantity": before}, after={"quantity": after, "reason": reason},
          via={"ai": "ai", "import": "import", "sale": "system"}.get(source, "api"))
    db.flush()
    return event


def undo_event(db: Session, *, actor_id: uuid.UUID, business_id: uuid.UUID, event_id: uuid.UUID) -> InventoryEvent:
    """Revert a manual/AI/import adjustment with a compensating event. Only the latest event of a product is undoable."""
    ev = db.scalar(select(InventoryEvent).where(InventoryEvent.id == event_id, InventoryEvent.business_id == business_id))
    if ev is None:
        raise NotFound("Hadisə tapılmadı")
    if ev.reverted or ev.event_type in NOT_UNDOABLE:
        msg = ("Satış və qaytarma hərəkətləri burada geri alınmır: satışı «Satışlar» bölməsindən ləğv edin."
               if ev.event_type in ("sale", "return") else "Bu dəyişiklik geri qaytarıla bilməz")
        raise Conflict(msg, code="not_undoable")
    latest = db.scalar(
        select(InventoryEvent).where(InventoryEvent.product_id == ev.product_id)
        .order_by(InventoryEvent.created_at.desc(), InventoryEvent.id.desc()).limit(1)
    )
    if latest is None or latest.id != ev.id:
        raise Conflict("Yalnız məhsulun son stok dəyişikliyi geri qaytarıla bilər", code="not_latest")
    undo = adjust_stock(db, actor_id=actor_id, business_id=business_id, product_id=ev.product_id, delta=-ev.delta,
                        reason="Geri qaytarma", source="undo", event_type="undo")
    undo.reverts_event_id = ev.id
    ev.reverted = True
    db.flush()
    return undo


def list_events(db: Session, business_id: uuid.UUID, *, product_id: uuid.UUID | None = None,
                event_type: str | None = None, limit: int = 50, offset: int = 0) -> tuple[list[dict], int]:
    stmt = (
        select(InventoryEvent, Product.name, Product.unit)
        .join(Product, Product.id == InventoryEvent.product_id)
        .where(InventoryEvent.business_id == business_id)
    )
    if product_id:
        stmt = stmt.where(InventoryEvent.product_id == product_id)
    if event_type:
        stmt = stmt.where(InventoryEvent.event_type == event_type)
    total = db.scalar(select(func.count()).select_from(stmt.with_only_columns(InventoryEvent.id).subquery())) or 0
    rows = db.execute(stmt.order_by(InventoryEvent.created_at.desc(), InventoryEvent.id.desc())
                      .offset(max(offset, 0)).limit(min(limit, 200))).all()
    # latest event per product, only for the products on this page (window function, not a full ledger scan)
    pids = {ev.product_id for ev, _, _ in rows}
    latest: set[uuid.UUID] = set()
    if pids:
        ranked = (
            select(InventoryEvent.id, func.row_number().over(
                partition_by=InventoryEvent.product_id,
                order_by=(InventoryEvent.created_at.desc(), InventoryEvent.id.desc())).label("rn"))
            .where(InventoryEvent.business_id == business_id, InventoryEvent.product_id.in_(pids))
            .subquery()
        )
        latest = set(db.scalars(select(ranked.c.id).where(ranked.c.rn == 1)).all())
    items = [serialize_event(ev, name, unit, undoable=(ev.id in latest and not ev.reverted
                                                       and ev.event_type not in NOT_UNDOABLE))
             for ev, name, unit in rows]
    return items, total


def serialize_event(ev: InventoryEvent, product_name: str | None = None, unit: str | None = None, undoable: bool = False) -> dict:
    return {
        "id": str(ev.id), "product_id": str(ev.product_id), "product_name": product_name, "unit": unit,
        "event_type": ev.event_type, "source": ev.source, "delta": qty_str(ev.delta),
        "quantity_before": qty_str(ev.quantity_before), "quantity_after": qty_str(ev.quantity_after),
        "reason": ev.reason, "reverted": ev.reverted, "undoable": undoable,
        "sale_id": str(ev.sale_id) if ev.sale_id else None, "created_at": iso(ev.created_at),
    }


def summary(db: Session, business_id: uuid.UUID) -> dict[str, Any]:
    from app.services.products import cost_value, retail_value, stock_status

    products = db.scalars(
        select(Product).where(Product.business_id == business_id, Product.is_archived.is_(False))
        .options(selectinload(Product.inventory))
    ).all()
    total_units = Decimal(0)
    total_cost = Decimal(0)
    total_retail = Decimal(0)
    missing_cost = 0
    low: list[dict] = []
    out = 0
    for p in products:
        q = Decimal(p.inventory.quantity) if p.inventory else Decimal(0)
        total_units += q
        cv, rv = cost_value(p), retail_value(p)
        if cv is None:
            missing_cost += 1
        else:
            total_cost += cv
        total_retail += rv or Decimal(0)
        st = stock_status(q, p.low_stock_threshold)
        if st == "out_of_stock":
            out += 1
        if st in ("low", "out_of_stock"):
            low.append({"id": str(p.id), "name": p.name, "sku": p.sku, "quantity": qty_str(q), "unit": p.unit,
                        "threshold": p.low_stock_threshold, "status": st})
    low.sort(key=lambda x: Decimal(x["quantity"]))
    return {
        "total_products": len(products),
        "total_units": qty_str(total_units),
        "inventory_value": money_str(total_cost),
        "retail_value": money_str(total_retail),
        "products_without_cost": missing_cost,
        "low_stock_count": len(low),
        "out_of_stock_count": out,
        "low_stock": low[:20],
    }


# ---------------------------------------------------------------------------
# CSV / XLSX import & export
# ---------------------------------------------------------------------------
EXPORT_COLUMNS = ["sku", "name", "category", "quantity", "unit", "purchase_price", "selling_price",
                  "inventory_value", "stock_status", "low_stock_threshold", "description", "updated_at"]


def export_rows(db: Session, business_id: uuid.UUID) -> list[dict]:
    from app.services.products import list_products, serialize_product

    items, total = list_products(db, business_id, page_size=500, sort="name")
    page = 2
    while len(items) < total:
        batch, _ = list_products(db, business_id, page_size=500, page=page, sort="name")
        if not batch:
            break
        items.extend(batch)
        page += 1
    rows = []
    for p in items:
        s = serialize_product(p)
        rows.append({**{k: s.get(k) for k in EXPORT_COLUMNS}, "category": (s["category"] or {}).get("name", "")})
    return rows


def spreadsheet_safe(value):
    # Text cells must never become spreadsheet formulas on download.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def export_csv(db: Session, business_id: uuid.UUID) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS)
    w.writeheader()
    for r in export_rows(db, business_id):
        w.writerow({k: spreadsheet_safe(v) for k, v in r.items()})
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def export_xlsx(db: Session, business_id: uuid.UUID) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventar"
    ws.append(EXPORT_COLUMNS)
    for r in export_rows(db, business_id):
        row = []
        for c in EXPORT_COLUMNS:
            v = r.get(c)
            if c in ("quantity", "purchase_price", "selling_price", "inventory_value") and v is not None:
                v = float(v)
            row.append(spreadsheet_safe(v))
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# Column names people actually use, in Azerbaijani, English, Turkish and Russian. Matched after `_norm_header`
# (diacritics folded, punctuation and "(AZN)"-style suffixes dropped).
_HEADER_ALIASES = {
    "sku": ["sku", "kod", "code", "mehsul kodu", "mal kodu", "artikul", "artikel", "barkod", "barcode", "stok kodu",
            "urun kodu", "артикул", "код", "штрихкод"],
    "name": ["name", "product", "product name", "item", "ad", "adi", "mehsul", "mehsullar", "mehsul adi",
             "mehsulun adi", "mal", "mallar", "mal adi", "malin adi", "urun", "urun adi", "наименование", "название",
             "товар", "наименование товара"],
    "quantity": ["quantity", "qty", "stock", "miqdar", "miqdari", "say", "sayi", "eded", "stok", "qaliq", "anbar",
                 "anbarda", "movcud", "movcud say", "adet", "miktar", "количество", "кол во", "остаток"],
    "unit": ["unit", "vahid", "olcu", "olcu vahidi", "birim", "ед изм", "единица", "единица измерения"],
    "purchase_price": ["purchase price", "cost", "alis", "alis qiymeti", "alis qiymet", "maya", "maya deyeri",
                       "alis fiyati", "закупочная цена", "себестоимость"],
    "selling_price": ["selling price", "price", "sale price", "qiymet", "qiymeti", "satis", "satis qiymeti",
                      "satis qiymet", "satis fiyati", "fiyat", "цена", "цена продажи"],
    "category": ["category", "group", "kateqoriya", "kategoriya", "kategori", "qrup", "категория", "группа"],
    "description": ["description", "tesvir", "qeyd", "aciklama", "описание", "примечание"],
    "low_stock_threshold": ["low stock threshold", "minimum stok", "min stok", "minimum", "min qaliq"],
}
_NUMERIC_FIELDS = ("quantity", "purchase_price", "selling_price", "low_stock_threshold")
_TOTAL_ROWS = {"cemi", "toplam", "total", "yekun", "umumi", "итого", "всего"}


def _norm_header(value) -> str:
    from app.services.lexicon import fold

    text = re.sub(r"\(.*?\)|\[.*?\]", " ", str(value or ""))  # "Qiymət (AZN)" -> "qiymet"
    return re.sub(r"\s+", " ", re.sub(r"[_.,/:*#№-]", " ", fold(text))).strip()


_ALIAS_INDEX = {_norm_header(a): field for field, aliases in _HEADER_ALIASES.items() for a in aliases}


def _header_field(value) -> str | None:
    h = _norm_header(value)
    if not h:
        return None
    if h in _ALIAS_INDEX:
        return _ALIAS_INDEX[h]
    # "Satış qiyməti 2024", "Anbardakı say": the longest alias the header starts or ends with
    hits = [a for a in _ALIAS_INDEX if len(a) > 2 and (h.startswith(a + " ") or h.endswith(" " + a))]
    return _ALIAS_INDEX[max(hits, key=len)] if hits else None


def clean_number(value: str) -> str:
    """'1 250,50 AZN' -> '1250.50', '1.250,5' -> '1250.5'. Text without digits comes back unchanged so it still
    fails validation with a readable error."""
    s = str(value).strip()
    if not re.search(r"\d", s):
        return s
    s = re.sub(r"[^\d,.\-]", "", s)
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    return s


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _find_header(rows: list[list[str]]) -> tuple[int, dict[int, str]] | None:
    """The first row among the first 20 that names the product (name or SKU) plus at least one other known column."""
    for i, row in enumerate(rows[:20]):
        cols: dict[int, str] = {}
        for j, v in enumerate(row):
            f = _header_field(v)
            if f and f not in cols.values():
                cols[j] = f
        if ("name" in cols.values() or "sku" in cols.values()) and len(cols) >= 2:
            return i, cols
    return None


def _guess_columns(rows: list[list[str]]) -> dict[int, str]:
    """No recognisable header: the first mostly-text column is the name, the first numeric one after it the quantity."""
    sample = [r for r in rows[:50] if any(r)]
    width = max((len(r) for r in sample), default=0)

    def values(j: int) -> list[str]:
        return [r[j] for r in sample if j < len(r) and r[j]]

    def numeric(j: int) -> float:
        vals = values(j)
        return sum(bool(re.fullmatch(r"-?[\d\s.,]+", v)) for v in vals) / len(vals) if vals else 0

    def filled(j: int) -> bool:
        return len(values(j)) >= max(1, len(sample) // 2)

    name_col = next((j for j in range(width) if filled(j) and numeric(j) < 0.5), None)
    if name_col is None:
        return {}
    cols = {name_col: "name"}
    qty_col = next((j for j in range(name_col + 1, width) if filled(j) and numeric(j) >= 0.8), None)
    if qty_col is not None:
        cols[qty_col] = "quantity"
    return cols


def _sheets(filename: str, content: bytes) -> list[list[list[str]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ValidationFailed("Excel faylı oxuna bilmədi. Faylı Excel-də açıb «.xlsx» kimi yenidən yadda saxlayın.")
        sheets = [wb.active] + [ws for ws in wb.worksheets if ws is not wb.active]
        return [[[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)] for ws in sheets if ws is not None]
    if name.endswith(".xls"):
        raise ValidationFailed("Köhnə .xls formatı dəstəklənmir. Faylı Excel-də açıb «.xlsx» kimi yadda saxlayın.")
    if name.endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        # the delimiter that splits the most lines; csv.Sniffer gets confused by a title line above the table
        lines = text.splitlines()[:50]
        delimiter = max(";,\t", key=lambda d: (sum(d in ln for ln in lines), sum(ln.count(d) for ln in lines)))
        return [[[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)]]
    raise ValidationFailed("Yalnız .xlsx və .csv faylları dəstəklənir")


def parse_upload(filename: str, content: bytes) -> list[dict[str, str]]:
    return read_upload(filename, content)["rows"]


def read_upload(filename: str, content: bytes) -> dict[str, Any]:
    """Finds the product table in a spreadsheet. The header may sit under a title block, on any sheet, in any of the
    languages above, or be missing altogether. Returns the mapped rows plus which file column became which field."""
    sheets = _sheets(filename, content)
    found = next(((rows, *hit) for rows in sheets if (hit := _find_header(rows))), None)
    guessed = found is None
    if found is None:
        rows = next((r for r in sheets if any(any(c) for c in r)), [])
        cols = _guess_columns(rows)
        if not cols:
            raise ValidationFailed("Faylda məhsul cədvəli tapılmadı. Birinci sətirdə «Məhsul», «Miqdar», «Qiymət» "
                                   "kimi başlıqlar olsun.")
        # treat the first row as a header unless its quantity cell already holds a number
        q = next((j for j, f in cols.items() if f == "quantity"), None)
        first = rows[0] if rows else []
        header_at = 0 if q is not None and q < len(first) and not re.search(r"\d", first[q]) else -1
        found = (rows, header_at, cols)
    rows, header_at, cols = found
    header = rows[header_at] if header_at >= 0 else []
    out = []
    for i in range(header_at + 1, len(rows)):
        r = rows[i]
        mapped = {f: (r[j] if j < len(r) else "") for j, f in cols.items()}
        if not any(mapped.values()) or _norm_header(mapped.get("name")) in _TOTAL_ROWS:
            continue  # blank lines and a "Cəmi" row under the table
        for f in _NUMERIC_FIELDS:
            if mapped.get(f):
                mapped[f] = clean_number(mapped[f])
        mapped["_row"] = i + 1
        out.append(mapped)
    if len(out) > 2000:
        raise ValidationFailed("Bir dəfəyə ən çox 2000 sətir idxal edilə bilər")
    columns = [{"column": header[j] if j < len(header) and header[j] else f"{j + 1}-ci sütun", "field": f}
               for j, f in sorted(cols.items())]
    return {"rows": out, "columns": columns, "ignored": [h for j, h in enumerate(header) if h and j not in cols],
            "guessed": guessed}


def import_rows(db: Session, *, actor_id: uuid.UUID, business_id: uuid.UUID, rows: list[dict[str, str]],
                dry_run: bool) -> dict[str, Any]:
    """Upsert by SKU (or exact name). Existing products get their quantity *set* via a ledger event."""
    from app.services.categories import get_or_create
    from app.services.products import create_product, update_product

    results = []
    created = updated = errors = 0
    for idx, r in enumerate(rows, start=2):
        idx = r.get("_row", idx)
        line = {"row": idx, "name": r.get("name") or r.get("sku"), "sku": r.get("sku") or None,
                "quantity": r.get("quantity") or None, "unit": r.get("unit") or None,
                "purchase_price": r.get("purchase_price") or None, "selling_price": r.get("selling_price") or None,
                "category": r.get("category") or None, "old_quantity": None}
        try:
            if not r.get("name") and not r.get("sku"):
                raise ValidationFailed("Ad və ya SKU tələb olunur")
            product = None
            if r.get("sku"):
                product = db.scalar(select(Product).where(Product.business_id == business_id, Product.sku == r["sku"]))
            if product is None and r.get("name"):
                product = db.scalar(select(Product).where(Product.business_id == business_id,
                                                          func.lower(Product.name) == r["name"].lower()))
            fields: dict[str, Any] = {}
            for f in ("name", "unit", "purchase_price", "selling_price", "description", "low_stock_threshold"):
                if r.get(f):
                    fields[f] = r[f]
            if "unit" in fields:
                from app.services.lexicon import parse_unit

                fields["unit"] = parse_unit(fields["unit"]) or fields["unit"]
            q = qty(r["quantity"]) if r.get("quantity") else None
            if q is not None and q < 0:
                raise ValidationFailed("Miqdar mənfi ola bilməz")
            with db.begin_nested():
                if r.get("category"):
                    fields["category_id"] = get_or_create(db, actor_id=actor_id, business_id=business_id,
                                                          name=r["category"], via="import").id
                if product is None:
                    if r.get("sku"):
                        fields["sku"] = r["sku"]
                    create_product(db, actor_id=actor_id, business_id=business_id, data=fields,
                                   initial_quantity=q or 0, via="import")
                    action = "create"
                    created += 1
                else:
                    line["old_quantity"] = qty_str(product.inventory.quantity) if product.inventory else None
                    line["name"] = r.get("name") or product.name
                    update_product(db, actor_id=actor_id, product=product, data=fields, via="import")
                    if q is not None and product.inventory and Decimal(product.inventory.quantity) != q:
                        adjust_stock(db, actor_id=actor_id, business_id=business_id, product_id=product.id,
                                     set_to=q, source="import", event_type="import", reason="Fayldan idxal")
                    action = "update"
                    updated += 1
            if q is not None:
                line["quantity"] = qty_str(q)
            results.append({**line, "status": "ok", "action": action})
        except Exception as exc:  # report per-row problems instead of failing the whole file
            errors += 1
            msg = getattr(exc, "message", None) or str(exc)
            results.append({**line, "status": "error", "error": msg})
    if dry_run:
        db.rollback()
    return {"dry_run": dry_run, "created": created, "updated": updated, "errors": errors, "rows": results}
